# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content as C

FONT = "Times New Roman"
OUT = os.path.join(C.SC, "MTN_Enterprise_Analytics_Strategy.docx")
doc = Document()

normal = doc.styles["Normal"]
normal.font.name = FONT; normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def sr(r, size=12, bold=False, italic=False):
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

def P(text, size=12, bold=False, italic=False, align="justify", sa=6, sb=0, indent=None, line15=True):
    p = doc.add_paragraph()
    p.alignment = {"justify":WD_ALIGN_PARAGRAPH.JUSTIFY,"left":WD_ALIGN_PARAGRAPH.LEFT,
                   "center":WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if line15 else WD_LINE_SPACING.SINGLE
    if indent: p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(C.clean(text)); sr(r, size, bold, italic)
    return p

# cover
doc.add_paragraph(); doc.add_paragraph()
for txt, size, bold in C.COVER:
    if txt == "__gap__":
        doc.add_paragraph()
        continue
    P(txt, size=size, bold=bold, align="center", sa=3)
doc.add_page_break()

def add_pagenum(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); sr(r, size=10)
    for tag, typ in [("w:fldChar","begin"),(None,None),("w:fldChar","end")]:
        pass
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"),"begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"),"preserve"); i.text="PAGE"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"),"end")
    r._r.append(b); r._r.append(i); r._r.append(e)

for el in C.DOC:
    k = el["k"]
    if k == "h1":
        p = P(el["t"], size=13, bold=True, align="left", sa=6, sb=16)
        p.paragraph_format.keep_with_next = True
    elif k == "h2":
        p = P(el["t"], size=12, bold=True, align="left", sa=4, sb=10)
        p.paragraph_format.keep_with_next = True
    elif k == "p":
        P(el["t"], size=el.get("size",12), indent=(0.3 if el.get("indent") else None))
    elif k == "pagebreak":
        doc.add_page_break()
    elif k == "fig":
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_before = Pt(8); pp.paragraph_format.space_after = Pt(2)
        pp.add_run().add_picture(el["path"], width=Inches(el["w"]))
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.LEFT
        c.paragraph_format.space_after = Pt(10); c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r1 = c.add_run(el["num"]+" "); sr(r1, 10.5, bold=True)
        r2 = c.add_run(C.clean(el["cap"])); sr(r2, 10.5, italic=True)
    elif k == "table":
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.LEFT
        c.paragraph_format.space_before = Pt(8); c.paragraph_format.space_after = Pt(2)
        c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r1 = c.add_run(el["num"]+" "); sr(r1,10.5,bold=True)
        r2 = c.add_run(C.clean(el["cap"])); sr(r2,10.5,italic=True)
        t = doc.add_table(rows=1, cols=len(el["headers"])); t.style="Table Grid"; t.alignment=1
        for i,h in enumerate(el["headers"]):
            cell = t.rows[0].cells[i]; cell.text=""
            pr = cell.paragraphs[0]; pr.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
            rr = pr.add_run(C.clean(h)); sr(rr,10.5,bold=True)
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),"EBEBEB")
            cell._tc.get_or_add_tcPr().append(shd)
        for row in el["rows"]:
            cells = t.add_row().cells
            for i,v in enumerate(row):
                cells[i].text=""; pr=cells[i].paragraphs[0]
                pr.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
                rr=pr.add_run(C.clean(v)); sr(rr,10.5)
        for i,w in enumerate(el["widths"]):
            for row in t.rows: row.cells[i].width=Inches(w)
        doc.add_paragraph().paragraph_format.space_after=Pt(4)
    elif k == "refs":
        for ref in el["items"]:
            p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.left_indent=Inches(0.5)
            p.paragraph_format.first_line_indent=Inches(-0.5)
            rr=p.add_run(C.clean(ref)); sr(rr,12)

for s in doc.sections:
    add_pagenum(s)
    s.top_margin=Inches(1.0); s.bottom_margin=Inches(1.0)
    s.left_margin=Inches(1.0); s.right_margin=Inches(1.0)

doc.save(OUT); print("docx saved:", OUT)
